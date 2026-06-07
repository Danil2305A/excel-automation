from dataclasses import dataclass

import openpyxl


@dataclass
class Practice:
    type: str
    name: str
    course: str

    sheet_name = "Практики"

    def __str__(self) -> str:
        return f"{self.type} - {self.name}"


@dataclass
class Specialization:
    code: str
    name: str
    profile: str

    sheet_name = "Титул"

    def __str__(self) -> str:
        return f"{self.code} {self.name} ({self.profile})"


class ExcelParser:
    def __init__(self, excel_filepath: str):
        self.excel_filepath = excel_filepath
        self.wb = openpyxl.load_workbook(excel_filepath)

    def get_specialization(self) -> Specialization:
        sheet = self.wb[Specialization.sheet_name]

        rows = list(sheet.iter_rows(values_only=True))
        row_with_code_and_name = [v for v in rows[28] if v is not None]
        row_with_profile = [v for v in rows[29] if v is not None]

        parts = row_with_code_and_name[1].strip().split(maxsplit=1)  # type: ignore
        spec_code = parts[0]
        spec_name = parts[1]
        spec_profile = row_with_profile[2].strip()  # type: ignore

        return Specialization(code=spec_code, name=spec_name, profile=spec_profile)

    def get_range_for_practices(self) -> dict[str, int]:
        """
        Возвращает диапазон обработки ячеек для листа "Практики":\n
        минимальный и максимальный номер строки,\n
        минимальный и максимальный номер столбца\n
        (нумерация с 1)
        """
        sheet = self.wb[Practice.sheet_name]

        header = list(sheet.iter_rows(values_only=True))[0]
        min_col = header.index("Название практики") + 1
        max_col = header.index("Курс") + 1

        first_col = list(sheet.iter_cols(values_only=True))[0]
        for i, c in enumerate(first_col):
            if c is not None and "Вид практики" in c:  # type: ignore
                min_row = i + 1
                break

        second_col = list(sheet.iter_cols(values_only=True))[1]
        for i, c in enumerate(second_col[::-1]):
            if c is not None:
                max_row = len(second_col) - i
                break

        return {
            "min_row": min_row,
            "max_row": max_row,
            "min_col": min_col,
            "max_col": max_col,
        }

    def get_practices(self) -> list[Practice]:
        range = self.get_range_for_practices()
        sheet = self.wb[Practice.sheet_name]

        practices: list[Practice] = []
        current_practice_type: str = ""

        for row in sheet.iter_rows(
            min_row=range["min_row"],
            max_row=range["max_row"],
            min_col=range["min_col"],
            max_col=range["max_col"],
            values_only=True,
        ):
            practice_type = row[range["min_col"] - 1]
            practice_name = row[range["min_col"]]
            practice_course = str(row[range["max_col"] - 1])

            if practice_type is not None and str(practice_type):
                current_practice_type = practice_type.split(":")[1]  # type: ignore

            if practice_name is not None and practice_course is not None:
                practices.append(
                    Practice(
                        type=current_practice_type.strip(),
                        name=practice_name.strip(),  # type: ignore
                        course=practice_course,
                    )
                )

        return practices
